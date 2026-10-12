"""Reports and downloads: the numbers must mean what the page says they mean."""

import csv
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from conftest import at
from convene import exports, reports


@pytest.fixture
def finished(svc, clock, people, workshop, attendees):
    """The 3-seat workshop: 5 sign up, 1 cancels (someone promoted), 2 of 3 turn up, 1 left waiting."""
    regs = [svc.register(g, workshop) for g in attendees(5)]
    svc.cancel_registration(regs[0].user_id, regs[0].id)          # guest 4 promoted
    clock.set(at(10, 9, 20))
    svc.check_in(people["staff"], regs[1].ticket_code)
    clock.set(at(10, 10, 5))
    svc.check_in(people["staff"], regs[3].ticket_code)
    clock.set(at(10, 13))                                           # after the end
    return regs


def test_event_report_numbers(svc, workshop, finished):
    r = reports.event_report(svc, workshop)
    assert (r.capacity, r.confirmed, r.checked_in, r.waitlisted, r.cancelled, r.promoted) == (3, 3, 2, 1, 1, 1)
    assert r.no_shows == 1
    assert r.attendance_rate == pytest.approx(2 / 3)
    assert r.fill_rate == 1


def test_no_shows_are_not_counted_before_the_end(svc, clock, workshop, finished):
    clock.set(at(10, 11))
    r = reports.event_report(svc, workshop)
    assert not r.finished and r.no_shows is None


def test_arrivals_in_15_minute_steps_from_doors_opening(svc, workshop, finished):
    buckets = reports.event_report(svc, workshop).arrivals
    assert buckets[0][0] == at(10, 9, 0)                            # doors open an hour before 10:00
    assert [n for _, n in buckets] == [0, 1, 0, 0, 1]               # 09:20 and 10:05
    assert all(b[0] - a[0] == timedelta(minutes=15) for a, b in zip(buckets, buckets[1:]))


def test_overview_counts_finished_events_only(svc, clock, people, room, hall, workshop, finished, attendees):
    upcoming = svc.create_event(people["organiser"], hall, "Next month", at(28, 10), at(28, 12), capacity=1)
    svc.publish(people["organiser"], upcoming.id)
    for g in attendees(3):
        svc.register(g, upcoming.id)                                # 1 confirmed, 2 waiting, not finished
    ov = reports.overview(svc, people["organiser"])
    assert [r.event.title for r in ov.reports] == ["Next month", "Phishing Awareness Workshop"]
    assert (ov.total_confirmed, ov.total_checked_in) == (3, 2)
    assert ov.turned_away == 1                                      # the upcoming waitlist doesn't count
    assert [r.event.title for r in ov.unmet_demand] == ["Next month", "Phishing Awareness Workshop"]
    assert [(v.name, v.events, v.no_show_rate) for v in ov.venues] == [("Training Room 2", 1, pytest.approx(1 / 3))]


def test_overview_only_shows_my_events(svc, people, workshop, finished):
    assert reports.overview(svc, people["other_organiser"]).reports == []


def test_attendee_csv(svc, workshop, finished):
    people = [r for r in svc.attendees(workshop) if r.status != "cancelled"]
    text = exports.attendees_csv(people, svc.waitlist_position)
    assert text.startswith("﻿")                                # Excel reads accents correctly
    rows = list(csv.DictReader(io.StringIO(text.lstrip("﻿"))))
    assert [r["Status"] for r in rows].count("Confirmed") == 3
    assert [r["Waitlist position"] for r in rows if r["Status"] == "Waitlisted"] == ["1"]
    assert sum(1 for r in rows if r["Checked in"]) == 2


def test_csv_and_excel_injection_is_neutralised(svc, workshop):
    evil = svc.create_user("evil@example.test", '=HYPERLINK("http://evil.example","x")', "attendee",
                           "correct-horse-battery")
    svc.register(evil.id, workshop)
    people = svc.attendees(workshop)
    assert "'=HYPERLINK" in exports.attendees_csv(people, svc.waitlist_position)
    wb = load_workbook(io.BytesIO(exports.attendees_xlsx(svc.event(workshop), people, svc.waitlist_position)))
    cells = [c for ws in wb for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(c.data_type == "s" for c in cells)


def test_attendee_workbook_has_a_sign_in_sheet_of_confirmed_people(svc, workshop, finished):
    people = [r for r in svc.attendees(workshop) if r.status != "cancelled"]
    wb = load_workbook(io.BytesIO(exports.attendees_xlsx(svc.event(workshop), people, svc.waitlist_position)))
    assert wb.sheetnames == ["Attendees", "Sign-in sheet"]
    names = [c.value for c in wb["Sign-in sheet"]["B"][5:] if c.value]
    assert names == sorted(names) and len(names) == 3               # confirmed only, alphabetical
    assert wb["Sign-in sheet"].print_title_rows == "$5:$5"


def test_overview_workbook_formulas(svc, people, workshop, finished):
    wb = load_workbook(io.BytesIO(exports.overview_xlsx("Andrea", reports.overview(svc, people["organiser"]))))
    ws = wb["Events"]
    assert ws["I4"].value == '=IF(F4=0,"",G4/F4)'                  # attendance is a live formula
    assert (ws["F4"].value, ws["G4"].value) == (3, 2)
    assert wb["By venue"]["E2"].value.startswith("=IF(")
