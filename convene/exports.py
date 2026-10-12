"""Attendee lists and reports as CSV and Excel files.

Names and emails come from people signing up, so they are untrusted. A name
like  =HYPERLINK("http://evil",...)  would run as a formula when a CSV is
opened in Excel ("CSV injection"). So:
  * in CSV files, text starting with = + - @ or a tab/return gets a leading '
    (the fix recommended by OWASP); Excel shows it as plain text
  * in .xlsx files, such text is stored as a text cell, never a formula
"""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

DANGEROUS_START = ("=", "+", "-", "@", "\t", "\r")
HEADER_FILL = PatternFill("solid", fgColor="1B1A2E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="999999")


def safe_csv(value):
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(DANGEROUS_START) else text


def _time(dt):
    return dt.strftime("%Y-%m-%d %H:%M") if dt else ""


def attendee_rows(attendees, waitlist_position):
    rows = []
    for r in attendees:
        rows.append({
            "Name": r.user.name, "Email": r.user.email, "Status": r.status.capitalize(),
            "Waitlist position": waitlist_position(r.id) or "",
            "Registered": _time(r.queued_at), "Checked in": _time(r.checked_in_at),
        })
    return rows


def attendees_csv(attendees, waitlist_position):
    buffer = io.StringIO()
    rows = attendee_rows(attendees, waitlist_position)
    writer = csv.DictWriter(buffer, fieldnames=["Name", "Email", "Status", "Waitlist position",
                                                "Registered", "Checked in"])
    writer.writeheader()
    for row in rows:
        writer.writerow({k: safe_csv(v) for k, v in row.items()})
    return "﻿" + buffer.getvalue()          # the BOM makes Excel read accents (é, ñ) correctly


def attendees_xlsx(event, attendees, waitlist_position):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendees"
    rows = attendee_rows(attendees, waitlist_position)
    columns = ["Name", "Email", "Status", "Waitlist position", "Registered", "Checked in"]
    _title(ws, event, len(columns))
    _table(ws, 4, columns, [[r[c] for c in columns] for r in rows], [28, 32, 13, 17, 18, 18])
    _fit_landscape(ws)
    _sign_in_sheet(wb.create_sheet("Sign-in sheet"), event, [r for r in attendees if r.status == "confirmed"])
    return _bytes(wb)


def overview_xlsx(organiser_name, ov):
    wb = Workbook()
    ws = wb.active
    ws.title = "Events"
    put(ws, 1, 1, f"Attendance report: events run by {organiser_name}", Font(bold=True, size=14))
    columns = ["Date", "Event", "Venue", "Status", "Capacity", "Confirmed", "Checked in", "No-shows",
               "Attendance", "Fill", "Still waiting", "Promoted", "Cancellations"]
    data = []
    for r in ov.reports:
        data.append([r.event.starts_at.strftime("%Y-%m-%d %H:%M"), r.event.title, r.event.venue.name,
                     "Finished" if r.finished else r.event.status.capitalize(), r.capacity, r.confirmed,
                     r.checked_in, r.no_shows, None, None, r.waitlisted, r.promoted, r.cancelled])
    _table(ws, 3, columns, data, [17, 34, 28, 11, 10, 11, 11, 10, 12, 8, 13, 10, 13])
    _fit_landscape(ws)
    for i, r in enumerate(ov.reports, start=4):
        # Rates are formulas, so they stay right if someone edits a count in Excel
        ws.cell(row=i, column=9, value=f'=IF(F{i}=0,"",G{i}/F{i})').number_format = "0%"
        ws.cell(row=i, column=10, value=f'=IF(E{i}=0,"",F{i}/E{i})').number_format = "0%"
        if not r.finished:
            ws.cell(row=i, column=9).value = None     # (passing value=None to cell() would leave the formula)

    vs = wb.create_sheet("By venue")
    _table(vs, 1, ["Venue", "Finished events", "Confirmed", "Checked in", "No-show rate"],
           [[v.name, v.events, v.confirmed, v.checked_in, None] for v in ov.venues], [30, 16, 12, 12, 14])
    for i in range(2, len(ov.venues) + 2):
        vs.cell(row=i, column=5, value=f'=IF(C{i}=0,"",(C{i}-D{i})/C{i})').number_format = "0%"
    return _bytes(wb)


# ---- Building blocks ------------------------------------------------------------------------

def put(ws, row, col, value, font=None):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str) and value.startswith(DANGEROUS_START):
        cell.data_type = "s"
        cell.quotePrefix = True
    if font:
        cell.font = font
    return cell


def _title(ws, event, width):
    put(ws, 1, 1, event.title, Font(bold=True, size=14))
    put(ws, 2, 1, f"{event.starts_at:%A %d %B %Y, %H:%M}-{event.ends_at:%H:%M} · {event.venue.name}",
        Font(color="6A6880"))


def _table(ws, top, columns, rows, widths):
    for col, (name, width) in enumerate(zip(columns, widths), start=1):
        cell = put(ws, top, col, name, HEADER_FONT)
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=top + 1):
        for c, value in enumerate(row, start=1):
            put(ws, r, c, value)
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    ws.auto_filter.ref = f"A{top}:{get_column_letter(len(columns))}{top + max(len(rows), 1)}"


def _sign_in_sheet(ws, event, confirmed):
    """A printable list for the door: tick or sign next to each name."""
    _title(ws, event, 4)
    put(ws, 3, 1, f"{len(confirmed)} confirmed · print and sign on arrival", Font(color="6A6880"))
    headers = ["#", "Name", "Arrived (time)", "Signature"]
    for col, (name, width) in enumerate(zip(headers, [5, 34, 16, 34]), start=1):
        cell = put(ws, 5, col, name, Font(bold=True))
        cell.border = Border(bottom=Side(style="medium"))
        ws.column_dimensions[get_column_letter(col)].width = width
    for i, r in enumerate(sorted(confirmed, key=lambda r: r.user.name.lower()), start=1):
        row = 5 + i
        put(ws, row, 1, i).alignment = Alignment(horizontal="left")
        put(ws, row, 2, r.user.name)
        put(ws, row, 3, r.checked_in_at.strftime("%H:%M") if r.checked_in_at else "")
        for col in range(1, 5):
            ws.cell(row=row, column=col).border = Border(bottom=THIN)
            ws.cell(row=row, column=col).alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 26                     # room to sign
    ws.print_title_rows = "5:5"                                # repeat the header on every printed page
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)


def _fit_landscape(ws):
    """Print one page wide, as many pages long as needed."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
